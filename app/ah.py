import io

import tabula
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font


def convert_ah_pdf_to_excel(pdf: io.BytesIO) -> io.BytesIO:
    """
    Convert AH receipt in PDF to excel, such that you can devide cost
    """

    def append(ws, values, font = None):
        """My version of openpyxl append with font"""
        ws.append(values)
        if font is not None:
            for col in range(1, len(values) + 1):
                cell = ws.cell(ws.max_row, col)
                cell.font = font

    # TODO: the whole flow with temporary csv is probably not necessary
    # It does help with pagination though (it combines all in 1)
    csv_file='temp.csv'
    tabula.convert_into(pdf, csv_file, output_format="csv", pages='all')

    header_block = [0, None]
    groceries_block = [None, None]

    with open(csv_file) as f:
        for row_i, line in enumerate(f.read().splitlines()):
            if row_i != 0 and line.startswith('Omschrijving'):
                # This is where the main table starts
                groceries_block[0] = row_i
                header_block[1] = row_i-1

            if line.startswith('Overig'):
                groceries_block[1] = row_i - 1


    groceries_df = pd.read_csv(csv_file, skiprows = lambda row_i: row_i < groceries_block[0] or row_i > groceries_block[1])
    header_df = pd.read_csv(csv_file, skiprows = lambda row_i: row_i < header_block[0] or row_i > header_block[1])

    wb = Workbook()
    bold_font = Font(bold=True)

    if 'Sheet' in wb.sheetnames:
        wb.remove(wb['Sheet'])  # Remove default sheet
    ws = wb.create_sheet("AH boodschappen", 0)

    append(ws, ['', '', "Omschrijving", "Bedrag"], font=bold_font)

    start_row_groceries = 1+len(header_df)+3
    end_row_groceries = start_row_groceries + len(groceries_df)

    ## Header part
    ws.append(['', '', "Boodschappen D&D", f'=SUM(D{start_row_groceries}:D{end_row_groceries})-SUM(E{start_row_groceries}:E{end_row_groceries})'])
    ws.append(['', '', "Boodschappen Ank", f'=SUM(E{start_row_groceries}:E{end_row_groceries})'])
    # add remaining header items
    for row_index, row, in header_df.iterrows():
        if row['Omschrijving'].startswith('Boodschappen') or row['Btw'].startswith('Totaal'):
            continue

        # this can be done a lot more efficient
        cost = float(row['Inclusief btw'].replace(',', '.'))
        ws.append(['', '', row['Omschrijving'], cost])
    ws.append(['', '', "Totaal", f'=SUM(D2:D{ws.max_row})'])

    # empty row
    ws.append([''])

    append(ws, ['#', 'Ank', 'Omschrijving', 'Bedrag', 'Bedrag Ank'], font=bold_font)

    for row_index, row in groceries_df.iterrows():
        row_i = ws.max_row + 1
        cost = row['Inclusief btw']
        if isinstance(cost, str):
            cost = float(cost.replace(',', '.'))
        # this conversion can be done a lot more efficient
        content = [row['Aantal'], '', row['Omschrijving'], cost, f'=IF(D{row_i}="","",B{row_i}/A{row_i}*D{row_i})']
        ws.append(content)

    ws.column_dimensions['A'].width = 2
    ws.column_dimensions['B'].width = 4
    ws.column_dimensions['C'].width = 40

    virtual_spreadsheet = io.BytesIO()
    wb.save(virtual_spreadsheet)
    return virtual_spreadsheet
